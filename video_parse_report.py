# -*- coding: utf-8 -*-
"""Build a concise, evidence-led Word report for parsed videos.

The report intentionally groups sampling frames by detected shot. Raw per-frame
observations remain in the database/downloadable evidence package, but are not
dumped verbatim into the reader-facing Word document.
"""

from __future__ import annotations

import collections
import datetime
import io
import os
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from text_tools import normalize_segments, to_simplified


PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

NAVY = RGBColor(0x0B, 0x25, 0x45)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x66, 0x6B, 0x73)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x22, 0x25, 0x2A)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"


def _set_run_font(run, size=11, bold=None, color=INK, east_asia="Microsoft YaHei"):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr()
    fonts = run._element.rPr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run._element.rPr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def _style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("VIDEO ANALYSIS  ·  视频解析证据报告")
    _set_run_font(run, size=8.5, bold=True, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("第 ")
    _set_run_font(run, size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    run = footer.add_run(" 页")
    _set_run_font(run, size=8.5, color=MUTED)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGIN_DXA.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int], indent=TABLE_INDENT_DXA) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _remove_cell_default_text(cell) -> None:
    paragraph = cell.paragraphs[0]
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)


def _add_cell_text(cell, text: str, size=9.5, bold=False, color=INK, align=None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.18
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text or ""))
    _set_run_font(run, size=size, bold=bold, color=color)


def _add_label_detail(doc, rows: list[tuple[str, str]], widths=(1800, 7560), header=None):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    if header:
        cells = table.add_row().cells
        cells[0].merge(cells[1])
        _set_cell_shading(cells[0], LIGHT_BLUE)
        _add_cell_text(cells[0], header, size=10.5, bold=True, color=NAVY)
    for label, value in rows:
        cells = table.add_row().cells
        _set_cell_shading(cells[0], LIGHT_GRAY)
        _add_cell_text(cells[0], label, size=9.5, bold=True, color=DARK_BLUE)
        _add_cell_text(cells[1], value, size=9.5)
    _set_table_geometry(table, list(widths))
    return table


def _add_definition(doc, label: str, value: str, after=5) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(f"{label}：")
    _set_run_font(run, size=11, bold=True, color=DARK_BLUE)
    run = paragraph.add_run(value)
    _set_run_font(run, size=11, color=INK)


def _plain_markdown(text: str) -> str:
    value = str(text or "")
    value = re.sub(r"(?m)^#{1,6}\s*", "", value)
    value = value.replace("**", "").replace("__", "")
    value = re.sub(r"(?m)^[-*]\s+", "", value)
    value = value.replace("`", "")
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def _clip(text: str, limit=180) -> str:
    value = re.sub(r"\s+", " ", _plain_markdown(text)).strip()
    if len(value) <= limit:
        return value
    cut = value[:limit].rstrip("，；、。 ")
    return cut + "…"


def _report_section(report: str, heading: str) -> str:
    pattern = rf"(?ms)^(?:#+\s*)?【{re.escape(heading)}】[^\S\r\n]*\n?(.*?)(?=^(?:#+\s*)?【|\Z)"
    match = re.search(pattern, report or "")
    return match.group(1).strip() if match else ""


def _conclusion_blocks(report: str) -> list[tuple[str, str]]:
    section = _report_section(report, "解析结论")
    rows = []
    for block in re.split(r"\n\s*\n", section):
        cleaned = block.strip()
        if not cleaned or cleaned == "---":
            continue
        match = re.match(r"\*\*(.+?)\*\*[：:]\s*(.*)", cleaned, re.S)
        if match:
            rows.append((_plain_markdown(match.group(1)), _plain_markdown(match.group(2))))
        else:
            rows.append(("分析", _plain_markdown(cleaned)))
    return rows


def _markdown_rows(text: str, first_cell_pattern: str) -> list[list[str]]:
    rows = []
    for line in (text or "").splitlines():
        stripped = line.strip()
        if not stripped.startswith("|") or not stripped.endswith("|"):
            continue
        cells = [_plain_markdown(cell.strip()) for cell in stripped.strip("|").split("|")]
        if cells and re.search(first_cell_pattern, cells[0]):
            rows.append(cells)
    return rows


def _observation_map(observations: str) -> dict[int, list[str]]:
    out = {}
    for cells in _markdown_rows(observations, r"帧\s*\d{3}"):
        match = re.search(r"帧\s*(\d{3})", cells[0])
        if match and len(cells) >= 5:
            out[int(match.group(1))] = cells
    return out


def _shot_report_rows(report: str) -> dict[int, list[str]]:
    out = {}
    section = _report_section(report, "关键帧时间轴")
    for cells in _markdown_rows(section, r"镜头\s*\d+"):
        match = re.search(r"镜头\s*(\d+)", cells[0])
        if match:
            out[int(match.group(1))] = cells
    return out


# 这些字符串代表"模型没给出画面观察"，不是证据本身。
_NON_EVIDENCE = re.compile(
    r"视觉取证失败|视觉分析未执行|未配置 API key|本地模型未安装|没有返回视觉观察"
    r"|本组视觉取证超时|本帧视觉取证返回空内容|当前代表帧无视觉描述|无视觉描述"
)


def _has_visual_evidence(observations: str, report: str) -> bool:
    """判断这条视频到底有没有拿到画面证据（而不是只有失败占位串）。"""
    observed = _observation_map(observations)
    if observed:
        return any(row[1].strip() and not _NON_EVIDENCE.search(row[1]) for row in observed.values())
    if _shot_report_rows(report):
        return any(len(row) > 1 and row[1].strip() for row in _shot_report_rows(report).values())
    return False


def _time_seconds(value) -> float:
    try:
        return float(value or 0)
    except Exception:
        return 0.0


def _format_duration(seconds: float) -> str:
    seconds = max(0.0, seconds)
    minutes, remain = divmod(seconds, 60)
    return f"{int(minutes):02d}:{remain:05.2f}"


def _subtitle_text_for_range(segments: list[dict], start_sec: float, end_sec: float) -> str:
    lines = []
    for segment in segments or []:
        start = _time_seconds(segment.get("start_ms")) / 1000
        end = _time_seconds(segment.get("end_ms")) / 1000
        if end >= start_sec and start <= end_sec:
            text = (segment.get("text") or "").strip()
            if text and text not in lines:
                lines.append(text)
    return " / ".join(lines) if lines else "该时段未检测到可确认台词"


def _group_shots(manifest: list[dict]) -> list[dict]:
    groups = collections.OrderedDict()
    for entry in manifest or []:
        shot_id = int(entry.get("shot_id") or 0)
        groups.setdefault(shot_id, []).append(entry)
    result = []
    # 检测器的镜头编号会给没有采样到帧的镜头留空号（1,2,3,5…），
    # 直接沿用会让读者以为漏了镜头；这里按实际出现的镜头连续重编号，
    # 原始编号留在 source_shot_id 里用于对齐逐镜报告。
    for seq, (shot_id, entries) in enumerate(groups.items(), start=1):
        if not entries:
            continue
        representative = entries[len(entries) // 2]
        result.append({
            "shot_id": seq,
            "source_shot_id": shot_id,
            "entries": entries,
            "representative": representative,
            "start": _time_seconds(entries[0].get("shot_start_sec")),
            "end": _time_seconds(entries[-1].get("shot_end_sec")),
            "range": entries[0].get("shot_range") or "时间段未知",
        })
    return result


def _title_from_report(item_title: str, report: str) -> str:
    match = re.search(r"《([^》]{1,40})》", report or "")
    if match:
        return f"《{match.group(1)}》"
    return item_title or "未命名视频"


def _ocr_summary(observations: str) -> list[tuple[str, str]]:
    candidates = [x.strip() for x in re.findall(r"《([^》\n]{1,20})》", observations or "")]
    counts = collections.Counter(x for x in candidates if "看不清" not in x)
    if counts:
        primary, count = counts.most_common(1)[0]
        variants = [f"《{name}》×{n}" for name, n in counts.most_common() if name != primary]
        rows = [
            ("高频标题候选", f"《{primary}》在证据帧中出现 {count} 次；这是 OCR 高频结果，仍建议对照原帧复核。"),
            ("冲突识别", "、".join(variants) if variants else "未发现其他成组标题候选"),
        ]
    else:
        rows = [("标题候选", "没有得到足够清晰、可重复验证的标题文字")]

    markers = []
    for pattern in (
        r"--ar\s+\d+[:/]\d+",
        r"--stylize\s+\d+",
        r"\b(?:Midjourney|MJ|Seedance|Kling|Sora|ComfyUI)\b",
    ):
        markers.extend(re.findall(pattern, observations or "", flags=re.I))
    rows.append(("原片提示词/参数", "、".join(dict.fromkeys(markers)) if markers else "未发现可确认的原始提示词、模型参数或教程步骤"))
    rows.append(("小字与水印", "低清小字、水印和平台角标仅作为来源证据，不纳入内容结论；看不清处不补写。"))
    return rows


def _add_title_block(doc: Document, title: str, item: dict, raw: dict, duration: float) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(10)
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("抖音视频 · 证据驱动解析")
    _set_run_font(run, size=10, bold=True, color=BLUE)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(title)
    _set_run_font(run, size=25, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run("视频解析报告")
    _set_run_font(run, size=13, color=MUTED)

    meta_rows = [
        ("作者", item.get("author_name") or "未获取"),
        ("时长", f"约 {duration:.2f} 秒" if duration else "未获取"),
        ("字幕来源", raw.get("subtitle_source") or ("本地 ASR" if raw.get("subtitle_segments") else "未提取")),
        ("原链接", item.get("source_url") or raw.get("share_url") or "未记录"),
        ("生成时间", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    _add_label_detail(doc, meta_rows, widths=(1800, 7560))


def _add_metric_strip(doc: Document, duration: float, shots: list[dict], manifest: list[dict], segments: list[dict]) -> None:
    doc.add_paragraph()
    table = doc.add_table(rows=2, cols=4)
    values = [f"{duration:.1f}s" if duration else "—", str(len(shots)), str(len(manifest)), str(len(segments))]
    labels = ["视频时长", "检测镜头", "采样关键帧", "ASR 段落"]
    for idx in range(4):
        _set_cell_shading(table.cell(0, idx), LIGHT_BLUE)
        _add_cell_text(table.cell(0, idx), values[idx], size=15, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
        _add_cell_text(table.cell(1, idx), labels[idx], size=8.5, bold=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_table_geometry(table, [2340, 2340, 2340, 2340])


def _add_summary(doc: Document, report: str, duration: float, shots: list[dict], visual_ok: bool = True) -> None:
    doc.add_heading("一页读懂", level=1)
    blocks = _conclusion_blocks(report) if visual_ok else []
    if blocks:
        for label, value in blocks:
            _add_definition(doc, label, value)
    elif not visual_ok:
        _add_definition(
            doc,
            "内容概况",
            "本次没有取得画面视觉分析，报告只保留可核对的事实：视频时长、镜头切分、代表帧与 ASR 台词。"
            "画面描述、景别、光色、叙事作用等推断项一律留空，不用模板句填充。",
        )

    if shots and duration:
        longest = max(shots, key=lambda x: x["end"] - x["start"])
        pace = duration / max(1, len(shots))
        _add_definition(
            doc,
            "节奏结构",
            f"共 {len(shots)} 个镜头区间，平均每镜约 {pace:.1f} 秒；最长的是镜头 {longest['shot_id']}"
            f"（{longest['range']}，约 {longest['end'] - longest['start']:.1f} 秒）。",
        )

    table = doc.add_table(rows=1, cols=1)
    _set_cell_shading(table.cell(0, 0), CALLOUT)
    _add_cell_text(
        table.cell(0, 0),
        "阅读说明：正文按检测到的真实镜头区间归并，每个镜头只展示一张代表帧。完整采样帧仍保存在素材记录中，可单独下载复核。",
        size=9.5,
        bold=True,
        color=DARK_BLUE,
    )
    _set_table_geometry(table, [9360])


def _add_subtitles(doc: Document, segments: list[dict]) -> None:
    doc.add_heading("字幕与音频转文本", level=1)
    if not segments:
        doc.add_paragraph("未提取到可确认的人声或台词。")
        return
    p = doc.add_paragraph("以下为本地 ASR 草稿，保留原始时间码；错字、同音字和断句需要结合原视频人工复核。")
    p.paragraph_format.space_after = Pt(6)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    headers = ["时间码", "音频转文本（ASR 草稿）"]
    for idx, text in enumerate(headers):
        _set_cell_shading(table.cell(0, idx), LIGHT_BLUE)
        _add_cell_text(table.cell(0, idx), text, size=9.5, bold=True, color=NAVY)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for segment in segments:
        cells = table.add_row().cells
        timecode = f"{segment.get('start_timecode') or ''}\n—\n{segment.get('end_timecode') or ''}".strip()
        _add_cell_text(cells[0], timecode, size=8.5, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
        _add_cell_text(cells[1], segment.get("text") or "", size=10)
    _set_table_geometry(table, [2200, 7160])


def _add_ocr(doc: Document, observations: str, visual_ok: bool = True) -> None:
    doc.add_heading("画面文字、隐藏信息与 OCR 复核", level=1)
    rows = _observation_map(observations)
    texts = [(row[0], row[4]) for row in rows.values() if row[4].strip()]
    paragraph = doc.add_paragraph("以下为独立 OCR 转录，低清小字和同形字仍需对照原帧复核；不使用语言模型补写。")
    paragraph.paragraph_format.keep_with_next = True
    if texts:
        _add_label_detail(doc, texts, widths=(2100, 7260))
    else:
        doc.add_paragraph("本次没有取得可用 OCR 文字，具体状态见逐帧证据记录；保留原图供人工读取。")


def _shot_fields(group: dict, obs_map: dict[int, list[str]], shot_rows: dict[int, list[str]], segments: list[dict], visual_ok: bool = True):
    rep = group["representative"]
    frame_index = int(rep.get("index") or 0)
    obs = obs_map.get(frame_index) or []
    report_row = shot_rows.get(group["shot_id"]) or shot_rows.get(group.get("source_shot_id")) or []
    subtitle = _clip(_subtitle_text_for_range(segments, group["start"], group["end"]), 170)

    # 没有视觉证据时只出"能核对的事实"，不再用模板句把 8 个字段填满。
    if not visual_ok and not report_row and not obs:
        return [
            ("同期台词", subtitle),
            ("画面文字", "未做 OCR（本次视觉模型不可用）"),
        ]

    visible = report_row[1] if len(report_row) > 1 else (obs[1] if len(obs) > 1 else "")
    action = report_row[2] if len(report_row) > 2 else ""
    composition = report_row[3] if len(report_row) > 3 else (obs[2] if len(obs) > 2 else "")
    light = report_row[4] if len(report_row) > 4 else (obs[3] if len(obs) > 3 else "")
    ocr = report_row[5] if len(report_row) > 5 else (obs[4] if len(obs) > 4 else "")
    narrative = report_row[7] if len(report_row) > 7 else ""
    return [
        ("画面与人物", _clip(visible, 220)),
        ("动作 / 表情", _clip(action, 180)),
        ("景别 / 构图", _clip(composition, 170)),
        ("灯光 / 色彩", _clip(light, 160)),
        ("同期台词", subtitle),
        ("画面文字", _clip(ocr, 150)),
        ("叙事作用", _clip(narrative, 160)),
    ]


def _add_shot_cards(doc: Document, shots: list[dict], observations: str, report: str, segments: list[dict]) -> None:
    doc.add_heading("镜头级视频解析", level=1)
    visual_ok = _has_visual_evidence(observations, report)
    if visual_ok:
        doc.add_paragraph("按硬切检测区间归并相似帧；镜头内的连续动作以代表帧与 ASR 对齐，不将静帧推断冒充连续运动事实。")
    else:
        note = doc.add_table(rows=1, cols=1)
        _set_cell_shading(note.cell(0, 0), CALLOUT)
        _add_cell_text(
            note.cell(0, 0),
            "本次没有取得画面视觉分析（视觉模型不可用或未返回有效结果），"
            "因此不输出画面描述、景别、光色等无法核对的结论。"
            "下面只保留每个镜头的代表帧、时间码和同期台词；模型恢复后重新拉片可补齐。",
            size=9,
            bold=True,
            color=DARK_BLUE,
        )
        _set_table_geometry(note, [9360])
        doc.add_paragraph()

    obs_map = _observation_map(observations)
    shot_rows = _shot_report_rows(report)
    for group in shots:
        rep = group["representative"]
        heading = doc.add_paragraph()
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(5)
        heading.paragraph_format.keep_with_next = True
        run = heading.add_run(f"镜头 {group['shot_id']:02d}  ·  {group['range']}")
        _set_run_font(run, size=12, bold=True, color=BLUE)

        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        left, right = table.rows[0].cells
        _set_cell_shading(left, LIGHT_GRAY)
        image_path = rep.get("path") or ""
        if image_path and os.path.exists(image_path):
            p = left.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            try:
                p.add_run().add_picture(image_path, width=Inches(1.65))
            except Exception:
                _add_cell_text(left, "代表帧加载失败", size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            _add_cell_text(left, "无可用代表帧", size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
        cap = left.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(0)
        run = cap.add_run(f"帧{int(rep.get('index') or 0):03d}")
        _set_run_font(run, size=8.5, bold=True, color=DARK_BLUE)
        run.add_break()
        run = cap.add_run(rep.get("timecode") or "时间码未知")
        _set_run_font(run, size=8.5, bold=True, color=DARK_BLUE)

        _remove_cell_default_text(right)
        for label, value in _shot_fields(group, obs_map, shot_rows, segments, visual_ok):
            p = right.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.12
            run = p.add_run(f"{label}：")
            _set_run_font(run, size=9, bold=True, color=DARK_BLUE)
            run = p.add_run(value)
            _set_run_font(run, size=9, color=INK)
        _set_table_geometry(table, [2800, 6560])


def _add_director_analysis(doc: Document, shots: list[dict], observations: str, report: str, duration: float, visual_ok: bool = True) -> None:
    doc.add_heading("视频内容详细分析", level=1)
    detail = _report_section(report, "视频内容详细分析")
    if detail:
        for paragraph in detail.split("\n\n"):
            doc.add_paragraph(_plain_markdown(paragraph))
        doc.add_paragraph("以上为本地模型摘要；逐帧模型观察与独立 OCR 见下方，仍需对照原视频核对。")
        return

    # 没有画面证据时，这里以前会输出 6 段写死的"以人物中近景为主""低饱和基底"等结论，
    # 看起来像分析，实际与这条视频无关，是用户反馈"无用信息"的主要来源。
    if not visual_ok:
        _add_definition(
            doc,
            "本次缺什么",
            "叙事与人物、摄影与构图、灯光与色彩这些栏目需要画面证据支撑。"
            "本次视觉模型没有返回有效结果，因此不给出这几项结论，避免把猜测当成拉片。",
        )
        _add_definition(
            doc,
            "怎么补",
            "确认模型通道可用后，在记录上点「重新拉片/补拉片」重跑，报告会补上逐镜画面描述与 OCR。"
            "已下载的视频、音频和关键帧不受影响，可以直接复用，不会重复下载。",
        )
        return

    blocks = dict(_conclusion_blocks(report))
    story = blocks.get("发生的事") or blocks.get("视频类型与主体") or "详见各镜头卡与逐镜台词。"
    _add_definition(doc, "叙事与人物", story)

    # 只引用报告里真实给出的结论；没有就明说没有，不再套"双人中近景/低饱和基底"这类万能句。
    all_text = _plain_markdown(observations)
    composition = blocks.get("摄影与构图") or blocks.get("构图")
    if not composition:
        composition = "报告未单列构图结论，逐镜景别与构图见上方镜头卡。" if all_text else "报告未提供构图结论。"
    _add_definition(doc, "摄影与构图", composition)

    color = blocks.get("灯光与色彩") or blocks.get("光影色彩")
    if not color:
        color = "报告未单列光影结论，逐帧光色见上方镜头卡。" if all_text else "报告未提供光影结论。"
    _add_definition(doc, "灯光与色彩", color)

    if shots and duration:
        longest = max(shots, key=lambda x: x["end"] - x["start"])
        pacing = f"{duration:.2f} 秒内检测到 {len(shots)} 个镜头区间。镜头 {longest['shot_id']} 持续约 {longest['end'] - longest['start']:.2f} 秒，是全片最长段落。"
    else:
        pacing = "剪辑节奏缺少完整镜头边界数据，需要回看视频确认。"
    _add_definition(doc, "剪辑与节奏", pacing)

    _add_definition(doc, "声音与画面", "ASR 台词以时间码对齐到镜头区间，只能作为语义参考；说话人身份、背景音乐、环境声与口型精确同步需人工复核。")
    _add_definition(doc, "内容价值", blocks.get("核心信息价值") or blocks.get("内容价值") or "报告未单列内容价值判断。")


def _add_fact_boundary(doc: Document, raw: dict, manifest: list[dict]) -> None:
    doc.add_page_break()
    doc.add_heading("证据范围与人工复核清单", level=1)
    rows = [
        ("可直接确认", "关键帧中的人物、造型、姿态、景别、构图、光色、可辨文字，以及关键帧真实时间码。"),
        ("算法检测", f"检测到 {len(_group_shots(manifest))} 个镜头区间；硬切边界是算法结果，复杂遮挡或渐变转场可能需要人工校正。"),
        ("需要复核", "连续运镜、转场细节、说话人归属、精确口型、背景音乐、环境声、低清小字及 OCR 冲突。"),
        ("未纳入结论", "平台 UI、水印、账号名和看不清的字符；不会把复刻建议冒充原片提示词。"),
        ("完整证据", f"原始采样帧共 {len(manifest)} 张，保存在该素材记录的“下载图片/关键帧”中。"),
    ]
    _add_label_detail(doc, rows, widths=(1900, 7460))


def build_video_parse_report(item: dict, raw: dict) -> io.BytesIO:
    """Return a ready-to-send DOCX stream for parse-mode video reports."""
    manifest = [entry for entry in (raw.get("frame_manifest") or []) if isinstance(entry, dict)]
    shots = _group_shots(manifest)
    segments = normalize_segments([entry for entry in (raw.get("subtitle_segments") or []) if isinstance(entry, dict)])
    # 旧记录里存的是繁体 ASR，出报告时统一转简体（新记录在入库时已转过）。
    report = to_simplified(raw.get("visual_analysis") or "")
    observations = to_simplified(raw.get("visual_observations") or "")
    visual_ok = _has_visual_evidence(observations, report)
    duration = 0.0
    if shots:
        duration = max(group["end"] for group in shots)
    elif manifest:
        duration = max(_time_seconds(entry.get("time_sec")) for entry in manifest)

    doc = Document()
    _style_document(doc)
    title = _title_from_report(item.get("title") or "", report)
    _add_title_block(doc, title, item, raw, duration)
    _add_metric_strip(doc, duration, shots, manifest, segments)
    _add_summary(doc, report, duration, shots, visual_ok)
    _add_director_analysis(doc, shots, observations, report, duration, visual_ok)
    _add_subtitles(doc, segments)
    _add_ocr(doc, observations, visual_ok)
    _add_shot_cards(doc, shots, observations, report, segments)
    _add_fact_boundary(doc, raw, manifest)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream


def build_video_parse_report_file(item: dict, raw: dict, output_path: str) -> str:
    stream = build_video_parse_report(item, raw)
    Path(output_path).write_bytes(stream.getvalue())
    return output_path
